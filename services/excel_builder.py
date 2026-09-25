"""In-Memory 4-Tab Styled Excel Workbook Builder for Tech Events Finder using openpyxl."""

import io
import logging
from typing import List, Dict, Any
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.ai_extractor import TechEventRecord

logger = logging.getLogger("techyupdates.excel_builder")

# Sheet definitions with modern color palettes
TABS_CONFIG = [
    {
        "name": "AI, Cloud & Data Summits",
        "header_fill": "312E81",     # Deep Indigo
        "header_text": "FFFFFF",
        "sub_fill": "4338CA",
        "zebra_fill": "EEF2FF",      # Light Indigo Tint
        "border_color": "C7D2FE",
        "category_match": "AI, Cloud & Data Summits",
    },
    {
        "name": "Developer & Open Source Confs",
        "header_fill": "0F766E",     # Deep Teal
        "header_text": "FFFFFF",
        "sub_fill": "14B8A6",
        "zebra_fill": "F0FDFA",      # Light Teal Tint
        "border_color": "99F6E4",
        "category_match": "Developer & Open Source Confs",
    },
    {
        "name": "Student Tech Fests & Workshops",
        "header_fill": "065F46",     # Deep Forest Green
        "header_text": "FFFFFF",
        "sub_fill": "10B981",
        "zebra_fill": "ECFDF5",      # Light Emerald Tint
        "border_color": "A7F3D0",
        "category_match": "Student Tech Fests & Workshops",
    },
    {
        "name": "Meetups, Demo Days & CFPs",
        "header_fill": "9A3412",     # Deep Rust / Amber
        "header_text": "FFFFFF",
        "sub_fill": "F97316",
        "zebra_fill": "FFF7ED",      # Light Orange Tint
        "border_color": "FED7AA",
        "category_match": "Meetups, Demo Days & CFPs",
    },
]

HEADERS = [
    "Event Name",
    "Organizer / Source",
    "Event Type",
    "Domain & Track",
    "Mode",
    "Location / City",
    "Start Date",
    "End Date",
    "Registration / RSVP Deadline",
    "Date Announced / Posted",
    "Pricing & Perks",
    "Direct Registration Link",
]

COLUMN_WIDTHS = [38, 24, 16, 30, 14, 24, 16, 16, 24, 20, 26, 28]


def build_excel_workbook(records: List[TechEventRecord]) -> bytes:
    """Compile all classified tech events into a 4-tab styled Excel workbook in memory."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1F2937")
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single", bold=True)
    why_font = Font(name="Segoe UI", size=9, italic=True, color="4B5563")

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for tab_cfg in TABS_CONFIG:
        ws = wb.create_sheet(title=tab_cfg["name"])
        ws.views.sheetView[0].showGridLines = True

        # Filter records for this tab
        tab_records = [r for r in records if r.category == tab_cfg["category_match"]]

        # 1. Title Banner Row (Row 1)
        ws.merge_cells("A1:L1")
        banner_cell = ws["A1"]
        banner_cell.value = f"TechyUpdates - {tab_cfg['name'].upper()} RADAR  |  {datetime.now(timezone.utc).strftime('%d %B %Y')}  |  {len(tab_records)} Verified Events"
        banner_cell.font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
        banner_cell.fill = PatternFill(start_color=tab_cfg["header_fill"], end_color=tab_cfg["header_fill"], fill_type="solid")
        banner_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # 2. Header Row (Row 2)
        ws.row_dimensions[2].height = 28
        for col_idx, (header_text, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS), start=1):
            cell = ws.cell(row=2, column=col_idx, value=header_text)
            cell.font = header_font
            cell.fill = PatternFill(start_color=tab_cfg["sub_fill"], end_color=tab_cfg["sub_fill"], fill_type="solid")
            cell.alignment = center_align
            cell.border = thin_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # 3. Data Rows
        current_row = 3
        for idx, rec in enumerate(tab_records):
            ws.row_dimensions[current_row].height = 36
            row_fill_color = tab_cfg["zebra_fill"] if idx % 2 == 1 else "FFFFFF"
            row_fill = PatternFill(start_color=row_fill_color, end_color=row_fill_color, fill_type="solid")

            # Clean apply url for hyperlink
            clean_url = (rec.registration_url or "https://techyupdates.dev").replace('"', '""')
            hyperlink_formula = f'=HYPERLINK("{clean_url}", "Register / RSVP")'

            row_values = [
                rec.title,
                rec.platform_or_source,
                rec.event_type,
                rec.domain_track,
                rec.mode,
                rec.location,
                rec.start_date,
                rec.end_date,
                rec.registration_deadline,
                rec.date_posted,
                rec.pricing_ticket,
                hyperlink_formula,
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.fill = row_fill
                cell.border = thin_border

                if col_idx == 12:
                    cell.font = link_font
                    cell.alignment = center_align
                elif col_idx in (3, 5, 7, 8, 9, 10):
                    cell.font = data_font
                    cell.alignment = center_align
                else:
                    cell.font = data_font
                    cell.alignment = left_align

            current_row += 1

        # Fallback if no records in tab
        if not tab_records:
            ws.row_dimensions[3].height = 26
            ws.merge_cells("A3:L3")
            empty_cell = ws["A3"]
            empty_cell.value = "No active events matched for this specific category today. Check back tomorrow!"
            empty_cell.font = Font(name="Segoe UI", size=10, italic=True, color="6B7280")
            empty_cell.alignment = Alignment(horizontal="center", vertical="center")
            current_row += 1

        # Freeze headers and enable auto-filters
        ws.freeze_panes = "A3"
        if tab_records:
            ws.auto_filter.ref = f"A2:L{current_row - 1}"

    # Export workbook to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()

    logger.info("[ExcelBuilder] Generated workbook with 4 tabs, size: %d bytes.", len(excel_bytes))
    return excel_bytes
